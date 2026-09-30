"""
CV parser: turns a CV file (PDF, Word, plain text) into clean text for the matcher.

GDPR note: CVs are only read into memory. Nothing here writes CV text to disk,
so keep it that way when the peer trial starts.

Usage:
    from skillradar.cv_parser import read_cv
    text = read_cv("my_cv.pdf")
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}

# PDF text extraction often turns typographic ligatures and bullets into odd characters
LIGATURES = {"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl"}
BULLETS = "•●▪■◦‣∙·➢➤►▶✓✔★-–—"


def read_cv(path: str | Path) -> str:
    """Read a CV file and return cleaned text."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = _read_pdf(path)
    elif suffix == ".docx":
        raw = _read_docx(path)
    elif suffix in {".txt", ".md"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported CV format '{suffix}'. Use one of: {', '.join(sorted(SUPPORTED))}")

    text = clean_text(raw)
    if len(text) < 50:
        raise ValueError(
            f"Only {len(text)} characters of text found in {path.name}. "
            "Is it a scanned image? Export the CV as a text-based PDF or .docx instead."
        )
    return text


def _read_pdf(path: Path) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            # x_tolerance keeps words in two-column layouts from being glued together
            pages.append(page.extract_text(x_tolerance=1.5) or "")
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    import docx  # python-docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                  # skills are often put in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    """Undo the usual PDF extraction damage without changing what the words say."""
    text = unicodedata.normalize("NFKC", text)             # also splits most ligatures
    for lig, repl in LIGATURES.items():
        text = text.replace(lig, repl)
    text = text.replace("\u00ad", "")                       # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)            # "Kuber-\nnetes" -> "Kubernetes"
    text = re.sub(rf"^[ \t]*[{re.escape(BULLETS)}]+[ \t]*", "", text, flags=re.M)  # leading bullets
    text = re.sub(r"[ \t]+", " ", text)                     # collapse spaces
    text = re.sub(r"\n{3,}", "\n\n", text)                  # collapse blank lines
    return text.strip()
